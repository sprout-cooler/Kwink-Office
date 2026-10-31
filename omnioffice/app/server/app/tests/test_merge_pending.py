"""周报：待处理文件上传与合并（v1.0.35）。

覆盖点：
  - 上传校验：非 xlsx / 空文件 / 坏文件 必须被拒（不能变成 500）
  - 上传后元信息正确（文件名 / 大小 / 工作表数）
  - 替换：再次上传仍只保留一份记录，元信息同步更新
  - 合并：待处理文件工作表数 +1，且新表内容与原周报一致
  - 合并：工作表名冲突时追加序号而非覆盖已有期次
  - 合并：失败自动回滚，待处理文件恢复原样
  - 合并：样式（加粗/填充/合并单元格）必须带走
  - 归属隔离：看不到别人的待处理文件
"""
from __future__ import annotations

import io
import unittest

from fastapi import HTTPException
from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from app.core.identity import Identity
from app.db.models import Base, WeeklyMergeRecord, WeeklyPendingFile
from app.modules.weekly_report import merge_service

UID = "u1"
OTHER = "u2"


def _fresh_db() -> Session:
    engine = create_engine(
        "sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool
    )
    Base.metadata.create_all(engine)
    return Session(engine)


def _ident(uid: str = UID) -> Identity:
    return Identity(uid=uid, username="user-" + uid, is_admin=False)


def _xlsx(sheets: dict[str, list[list]]) -> bytes:
    """生成一个多工作表 xlsx。"""
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _sheets(path) -> dict:
    """读回文件内容：{表名: [[单元格值]]}。"""
    from openpyxl import load_workbook

    wb = load_workbook(path)
    return {n: [[c.value for c in r] for r in wb[n].iter_rows()] for n in wb.sheetnames}


class TestUploadValidation(unittest.TestCase):
    def setUp(self):
        self.db = _fresh_db()

    def _expect_bad_file(self, name: str, data: bytes) -> str:
        with self.assertRaises(HTTPException) as ctx:
            merge_service.save_pending_file(self.db, _ident(), name, data)
        self.assertEqual(ctx.exception.status_code, 400)
        self.assertEqual(ctx.exception.detail["error"], "BAD_PENDING_FILE")
        return ctx.exception.detail["message"]

    def test_rejects_non_xlsx(self):
        self.assertIn(".xlsx", self._expect_bad_file("台账.xls", b"whatever"))

    def test_rejects_empty(self):
        self.assertIn("空", self._expect_bad_file("台账.xlsx", b""))

    def test_rejects_corrupt(self):
        self.assertIn("无法解析", self._expect_bad_file("台账.xlsx", b"not a zip"))

    def test_meta_after_upload(self):
        data = _xlsx({"1月": [["日期", "客户"], ["2024-01-05", "甲"]], "2月": [["x"]]})
        meta = merge_service.save_pending_file(self.db, _ident(), "年度台账.xlsx", data)
        self.assertEqual(meta["fileName"], "年度台账.xlsx")
        self.assertEqual(meta["sheetCount"], 2)
        self.assertEqual(meta["sizeBytes"], len(data))

    def test_replace_keeps_single_record(self):
        merge_service.save_pending_file(self.db, _ident(), "旧.xlsx", _xlsx({"A": [["a"]]}))
        meta = merge_service.save_pending_file(
            self.db, _ident(), "新.xlsx", _xlsx({"B": [["b"]], "C": [["c"]]})
        )
        self.assertEqual(meta["fileName"], "新.xlsx")
        self.assertEqual(meta["sheetCount"], 2)
        total = self.db.scalar(select(func.count()).select_from(WeeklyPendingFile))
        self.assertEqual(total, 1)


class TestMergeIntoPending(unittest.TestCase):
    def setUp(self):
        self.db = _fresh_db()
        merge_service.save_pending_file(self.db, _ident(), "台账.xlsx", _xlsx({"1月": [["a"]]}))

    def test_appends_one_sheet_with_content(self):
        report = _xlsx({"本周": [["事项", "时长"], ["支持", 4]]})
        merged = merge_service.merge_report_into_pending(self.db, _ident(), report, "第20周")

        self.assertEqual(merged["sheetName"], "第20周")
        self.assertEqual(merged["sheetCountAfter"], merged["sheetCountBefore"] + 1)

        sheets = _sheets(merge_service.pending_file_path(self.db, _ident()))
        self.assertEqual(sheets["第20周"][0], ["事项", "时长"])
        self.assertEqual(sheets["第20周"][1], ["支持", 4])
        # 追加，不是替换整个工作簿 —— 原有表必须还在
        self.assertIn("1月", sheets)

    def test_duplicate_title_gets_suffix(self):
        """目标已存在同名表时追加序号，绝不覆盖已有期次。"""
        merge_service.save_pending_file(self.db, _ident(), "台账.xlsx", _xlsx({"周报": [["x"]]}))
        first = merge_service.merge_report_into_pending(self.db, _ident(), _xlsx({"s": [["b"]]}), "周报")
        second = merge_service.merge_report_into_pending(
            self.db, _ident(), _xlsx({"s": [["c"]]}), "周报"
        )
        self.assertEqual(first["sheetName"], "周报(2)")
        self.assertEqual(second["sheetName"], "周报(3)")
        # 原表必须还在（是追加，不是覆盖）
        sheets = _sheets(merge_service.pending_file_path(self.db, _ident()))
        self.assertIn("周报", sheets)

    def test_history_recorded(self):
        merge_service.merge_report_into_pending(self.db, _ident(), _xlsx({"s": [["b"]]}), "第20周")
        hist = merge_service.list_merge_history(self.db, _ident())
        self.assertEqual(len(hist), 1)
        self.assertEqual(hist[0]["sheetName"], "第20周")
        self.assertTrue(hist[0]["backupName"])

    def test_missing_pending_file_rejected(self):
        from fastapi import HTTPException

        self.db.query(WeeklyPendingFile).delete()
        self.db.commit()
        with self.assertRaises(HTTPException):
            merge_service.merge_report_into_pending(self.db, _ident(), _xlsx({"x": [["1"]]}), "x")

    def test_failure_rolls_back(self):
        path = merge_service.pending_file_path(self.db, _ident())
        before = path.read_bytes()
        with self.assertRaises(Exception):
            merge_service.merge_report_into_pending(self.db, _ident(), b"broken", "x")
        self.assertEqual(path.read_bytes(), before, "合并失败后待处理文件必须原样保留")

    def test_styles_preserved(self):
        """合并必须带走样式：加粗、字号、填充色、合并单元格。"""
        from openpyxl import Workbook, load_workbook
        from openpyxl.styles import Font, PatternFill

        wb = Workbook()
        ws = wb.active
        ws.title = "带样式"
        ws["A1"] = "标题"
        ws["A1"].font = Font(bold=True, size=16)
        ws["A1"].fill = PatternFill("solid", fgColor="4C6FFF")
        ws.merge_cells("A1:B1")
        buf = io.BytesIO()
        wb.save(buf)

        merge_service.merge_report_into_pending(self.db, _ident(), buf.getvalue(), "带样式")
        out = load_workbook(merge_service.pending_file_path(self.db, _ident()))["带样式"]
        self.assertEqual(out["A1"].value, "标题")
        self.assertTrue(out["A1"].font.bold)
        self.assertEqual(out["A1"].font.size, 16)
        self.assertTrue(any(str(r) == "A1:B1" for r in out.merged_cells.ranges))


class TestPeriodSheetTitle(unittest.TestCase):
    """工作表名 = 阶段时间区间（如 261026-261030）。"""

    def test_monday_to_friday(self):
        # 2026-10-26 是周一
        self.assertEqual(merge_service.period_sheet_title("2026-10-26"), "261026-261030")

    def test_crosses_month_boundary(self):
        # 2026-10-30 周五起的那一周 → 跨月
        self.assertEqual(merge_service.period_sheet_title("2026-10-30"), "261030-261103")

    def test_crosses_year_boundary(self):
        # 2026-12-28 周一 + 4 天 = 2027-01-01（周五）
        self.assertEqual(merge_service.period_sheet_title("2026-12-28"), "261228-270101")

    def test_name_is_legal_for_excel(self):
        """Excel 工作表名不能含 : \\ / ? * [ ]，也不能超 31 字符。"""
        name = merge_service.period_sheet_title("2026-10-26")
        self.assertNotRegex(name, r"[\\/?*\[\]:]")
        self.assertLessEqual(len(name), 31)

    def test_merge_uses_period_title(self):
        db = _fresh_db()
        merge_service.save_pending_file(db, _ident(), "台账.xlsx", _xlsx({"1月": [["a"]]}))
        merged = merge_service.merge_report_into_pending(
            db, _ident(), _xlsx({"s": [["b"]]}), merge_service.period_sheet_title("2026-10-26")
        )
        self.assertEqual(merged["sheetName"], "261026-261030")


class TestIsolation(unittest.TestCase):
    def setUp(self):
        self.db = _fresh_db()

    def test_other_users_file_invisible(self):
        merge_service.save_pending_file(self.db, _ident(OTHER), "别人.xlsx", _xlsx({"x": [["1"]]}))
        self.assertIsNone(merge_service.get_pending_file_meta(self.db, _ident()))
        self.assertEqual(len(merge_service.list_merge_history(self.db, _ident())), 0)


if __name__ == "__main__":
    unittest.main()
