"""周报合并：把导出生成的表格复制成「待处理文件」里的一个新工作表。

为什么需要它
------------
用户手上有自己的待处理台账 xlsx（可能带自己的格式与多个工作表）。每期周报
导出后，需要把它**追加**进这个台账形成历史累积，而不是每次拿到孤立文件
再手工粘贴。

关键实现约束（实测得出，别踩）
------------------------------
1. **不能用 `Workbook.copy_worksheet()`** —— 跨工作簿直接抛
   `ValueError: Cannot copy between worksheets from different workbooks`。

2. **也不能复制 `cell._style`** —— 那是 openpyxl 内部的 `StyleArray`，
   存的是**源 workbook 样式表的下标**。跨工作簿赋值后目标 workbook 的字体/
   填充表里没有对应下标 → 读取时 `IndexError`。
   **更危险的是保存时不报错、打开时才崩**，会静默产出坏文件。
   → 正确做法：逐属性 `copy(cell.font)` / `copy(cell.fill)` / ... ，
     让 openpyxl 自动把样式对象注册进目标 workbook。

3. **先备份再覆盖** —— 合并是就地替换。若中途失败（磁盘满、模板出错），
   原文件已被截断就不可逆。所以：备份 → 内存完成合并 → 落盘；失败自动回滚。
   备份只保留最近若干份，避免无限膨胀。
"""
from __future__ import annotations

import re
from copy import copy
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path

from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from ...core.config import settings
from ...core.identity import Identity
from ...db.models import WeeklyMergeRecord, WeeklyPendingFile


def _bad_file(message: str) -> HTTPException:
    """待处理文件相关的 400。

    刻意用 HTTPException 而不是 biz_dashboard 的 BizDashboardError ——
    后者只在业务大屏模块注册了异常处理器，在周报模块抛会冒泡成 500，
    把「文件格式不对」这种可自行修复的问题说成服务内部错误。
    """
    return HTTPException(
        status_code=400, detail={"error": "BAD_PENDING_FILE", "message": message}
    )


# Excel 工作表名约束：≤31 字符，且不能含 \ / ? * [ ] :
_ILLEGAL = re.compile(r"[\\/?*\[\]:]")
_SHEET_MAX = 31
# 每个用户保留的备份份数
KEEP_BACKUPS = 5
# 待处理文件大小上限（MB）
MAX_PENDING_MB = 10


def pending_root() -> Path:
    """待处理文件存放根目录（DATA_DIR 下，升级保留）。"""
    return Path(settings.data_dir) / "weekly-report" / "pending"


def _pending_file(uid: str) -> Path:
    return pending_root() / f"{_safe_uid(uid)}.xlsx"


def _backup_dir() -> Path:
    return pending_root() / "backups"


def _safe_uid(uid: str) -> str:
    """uid 参与文件名，必须挡住目录穿越。"""
    cleaned = re.sub(r"[^0-9A-Za-z_.-]", "_", uid or "anon")[:64]
    return cleaned or "anon"


def _sheet_title(raw: str) -> str:
    cleaned = _ILLEGAL.sub("_", (raw or "").strip())[:_SHEET_MAX]
    return cleaned or "周报"


def _unique_title(wb, base: str) -> str:
    """工作表名去重：同名追加 (2)(3)，**不覆盖**已有期次。"""
    if base not in wb.sheetnames:
        return base
    head = base[: _SHEET_MAX - 4]
    i = 2
    while f"{head}({i})" in wb.sheetnames:
        i += 1
    return f"{head}({i})"


def _copy_values_and_styles(src_ws, dst_ws) -> None:
    """逐格复制值与样式属性（不复制 `_style` 下标，见模块 docstring）。"""
    for r in range(1, src_ws.max_row + 1):
        if src_ws.row_dimensions[r].height:
            dst_ws.row_dimensions[r].height = src_ws.row_dimensions[r].height
        for c in range(1, src_ws.max_column + 1):
            s = src_ws.cell(row=r, column=c)
            t = dst_ws.cell(row=r, column=c, value=s.value)
            if not s.has_style:
                continue
            t.font = copy(s.font)
            t.fill = copy(s.fill)
            t.border = copy(s.border)
            t.alignment = copy(s.alignment)
            t.number_format = s.number_format


# ---------------------------------------------------------------- 元信息


def get_pending_file_meta(db: Session, identity: Identity) -> dict | None:
    """当前待处理文件的元信息；没有则返回 None。"""
    row = db.scalar(
        select(WeeklyPendingFile).where(WeeklyPendingFile.owner_uid == identity.uid)
    )
    if row is None:
        return None
    return {
        "fileName": row.file_name,
        "sizeBytes": row.size_bytes,
        "sheetCount": row.sheet_count,
        "updatedAt": row.updated_at.isoformat() if row.updated_at else "",
    }


def pending_file_path(db: Session, identity: Identity) -> Path | None:
    """待处理文件的实际路径；不存在则返回 None。"""
    row = db.scalar(
        select(WeeklyPendingFile).where(WeeklyPendingFile.owner_uid == identity.uid)
    )
    if row is None:
        return None
    p = Path(row.stored_path)
    return p if p.is_file() else None


def _validate_upload(file_name: str, data: bytes) -> None:
    """上传校验：扩展名 → 大小 → 能否被 openpyxl 解析（三步缺一不可）。"""
    if not (file_name or "").lower().endswith(".xlsx"):
        raise _bad_file("待处理文件仅支持 .xlsx 格式")
    if not data:
        raise _bad_file("文件内容为空")
    limit = MAX_PENDING_MB * 1024 * 1024
    if len(data) > limit:
        raise _bad_file(f"文件过大（{len(data) / 1024 / 1024:.1f}MB），上限 {MAX_PENDING_MB}MB")
    # 扩展名可以随便改，只有真正解析一次才能确认它是合法 xlsx
    from openpyxl import load_workbook

    try:
        probe = load_workbook(BytesIO(data), read_only=True)
        probe.close()
    except Exception as e:
        raise _bad_file(f"无法解析该 xlsx 文件：{type(e).__name__}") from None


def save_pending_file(db: Session, identity: Identity, file_name: str, data: bytes) -> dict:
    """上传待处理文件；已存在则**替换**——全库对每个用户只保留一份。"""
    _validate_upload(file_name, data)

    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(data))
    root = pending_root()
    root.mkdir(parents=True, exist_ok=True)
    target = _pending_file(identity.uid)

    # 原子写：先落临时文件再 replace，避免写一半断电留下半截 xlsx
    tmp = target.with_name(target.name + ".tmp")
    wb.save(tmp)
    tmp.replace(target)

    row = db.scalar(
        select(WeeklyPendingFile).where(WeeklyPendingFile.owner_uid == identity.uid)
    )
    if row is None:
        row = WeeklyPendingFile(owner_uid=identity.uid)
        db.add(row)
    row.file_name = (file_name or "").strip()[:255]
    row.stored_path = str(target)
    row.size_bytes = len(data)
    row.sheet_count = len(wb.sheetnames)
    db.commit()
    return get_pending_file_meta(db, identity) or {}


def delete_pending_file(db: Session, identity: Identity) -> dict:
    row = db.scalar(
        select(WeeklyPendingFile).where(WeeklyPendingFile.owner_uid == identity.uid)
    )
    if row is None:
        return {"ok": True, "removed": False}
    p = Path(row.stored_path)
    if p.is_file():
        p.unlink()
    db.delete(row)
    db.commit()
    return {"ok": True, "removed": True}


# ---------------------------------------------------------------- 合并


def merge_report_into_pending(
    db: Session,
    identity: Identity,
    xlsx_bytes: bytes,
    sheet_title_hint: str = "",
    batch_id: int = 0,
) -> dict:
    """执行导出后的合并：把周报表格追加成待处理文件的一个**新工作表**，就地替换原文件。

    ⚠️ 前置校验：
      - 待处理文件必须已上传（否则 400，给出可行动提示）；
      - 合并是**就地覆盖**，所以先备份；失败自动还原。

    ⚠️ 为什么只复制主工作表（第一个 sheet）：
      待处理文件是**累计台账**，一次合并对应一期周报，追加一张表最直观；
      源 xlsx 若有多表，全部塞进去会让台账结构失控。
    """
    current = pending_file_path(db, identity)
    if current is None:
        raise HTTPException(
            status_code=400,
            detail={
                "error": "NO_PENDING_FILE",
                "message": "还没有上传待处理文件，请先上传后再合并",
            },
        )

    # 先备份原文件：合并是就地覆盖，一旦中途失败原文件就被截断了
    backups = _backup_dir()
    backups.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now(timezone.utc).strftime("%Y%m%d-%H%M%S")
    backup_path = backups / f"{_safe_uid(identity.uid)}-{stamp}.xlsx"
    original = current.read_bytes()
    backup_path.write_bytes(original)

    try:
        source = load_workbook(BytesIO(xlsx_bytes))
        src_ws = source[source.sheetnames[0]]

        target = load_workbook(BytesIO(original))
        title = _unique_title(target, _sheet_title(sheet_title_hint or src_ws.title))

        dst_ws = target.create_sheet(title)
        dst_ws.sheet_view.showGridLines = src_ws.sheet_view.showGridLines
        _copy_values_and_styles(src_ws, dst_ws)
        for rng in list(src_ws.merged_cells.ranges):
            dst_ws.merge_cells(str(rng))
        for key, dim in src_ws.column_dimensions.items():
            if dim and dim.width:
                dst_ws.column_dimensions[key].width = dim.width

        target.save(current)
        after = len(load_workbook(current).sheetnames)
    except HTTPException:
        raise
    except Exception as e:
        # 回滚：合并失败就把备份写回去，绝不留半截文件给用户
        current.write_bytes(original)
        raise HTTPException(
            status_code=500,
            detail={
                "error": "MERGE_FAILED",
                "message": f"合并失败，已自动还原原文件：{type(e).__name__}",
            },
        ) from None

    row = db.scalar(
        select(WeeklyPendingFile).where(WeeklyPendingFile.owner_uid == identity.uid)
    )
    before = row.sheet_count if row else 0
    if row is not None:
        row.sheet_count = after
        row.size_bytes = current.stat().st_size

    record = WeeklyMergeRecord(
        owner_uid=identity.uid,
        batch_id=batch_id,
        sheet_name=title,
        sheet_count_before=before,
        sheet_count_after=after,
        backup_path=str(backup_path),
    )
    db.add(record)
    db.commit()
    _prune_backups(backups, identity.uid)

    return {
        "sheetName": title,
        "sheetCountBefore": before,
        "sheetCountAfter": after,
        "backupName": backup_path.name,
        "mergedAt": record.created_at.isoformat() if record.created_at else "",
    }


def _prune_backups(root: Path, uid: str) -> None:
    """只保留最近 KEEP_BACKUPS 份备份，其余删除（避免备份无限膨胀）。"""
    files = sorted(
        root.glob(f"{_safe_uid(uid)}-*.xlsx"), key=lambda p: p.stat().st_mtime, reverse=True
    )
    for old in files[KEEP_BACKUPS:]:
        old.unlink(missing_ok=True)


def list_merge_history(db: Session, identity: Identity, limit: int = 50) -> list[dict]:
    """合并历史（最近在前）。"""
    rows = db.scalars(
        select(WeeklyMergeRecord)
        .where(WeeklyMergeRecord.owner_uid == identity.uid)
        .order_by(WeeklyMergeRecord.id.desc())
        .limit(limit)
    ).all()
    return [
        {
            "id": r.id,
            "sheetName": r.sheet_name,
            "sheetCountBefore": r.sheet_count_before,
            "sheetCountAfter": r.sheet_count_after,
            "backupName": Path(r.backup_path).name if r.backup_path else "",
            "createdAt": r.created_at.isoformat() if r.created_at else "",
        }
        for r in rows
    ]
